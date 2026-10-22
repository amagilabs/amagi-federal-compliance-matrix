import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from amagi_federal_compliance_matrix import __version__
from amagi_federal_compliance_matrix.cli import main


def test_version(capsys):
    with pytest.raises(SystemExit) as e:
        main(["--version"])
    assert e.value.code == 0
    assert __version__ in capsys.readouterr().out


def test_no_args_prints_help(capsys):
    assert main([]) == 1
    assert "usage:" in capsys.readouterr().out


def test_sample_matrix(sample_pdf, tmp_path, capsys):
    assert main([str(sample_pdf), "-o", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "33 requirements: 15 offeror, 14 contractor, 4 government" in out
    assert "by UCF section: B 1, C 15, L 14, M 3" in out
    assert "5 FAR/DFARS clauses referenced (4 with titles)" in out
    xlsx = tmp_path / "sample_solicitation_matrix.xlsx"
    assert xlsx.exists() and (tmp_path / "sample_solicitation_matrix.md").exists()
    wb = load_workbook(xlsx)
    assert wb["Matrix"].max_row == 34
    clauses = [r[0].value for r in wb["Clauses"].iter_rows(min_row=2)]
    assert "52.204-21" in clauses and "52.212-1" in clauses


def test_sample_rows(sample_pdf, capsys):
    assert main(["matrix", str(sample_pdf), "--stdout", "json"]) == 0
    reqs = {r["id"]: r for r in json.loads(capsys.readouterr().out)["requirements"]}
    assert reqs["L.4-01"]["requirement"] == "The offeror shall describe its approach to each task in Section C."
    assert reqs["C.3.1-02"]["requirement"] == "The Contractor shall: (a) answer 90 percent of calls within 60 seconds;"
    assert reqs["C.3.2-03"]["requirement"].endswith("(c) a monthly asset report.")
    assert reqs["M.2-01"]["actor"] == "Government"
    assert reqs["C.3.3-01"]["kind"] == "Prohibition"
    assert not any("Business day" in r["requirement"] for r in reqs.values())
    assert not any("OER-26-R-0042" in r["requirement"] for r in reqs.values())


def test_options_before_file_and_formats(sample_pdf, tmp_path, capsys):
    assert main(["-o", str(tmp_path), "--name", "x", "--format", "csv,json", "-q", str(sample_pdf)]) == 0
    printed = sorted(Path(p).name for p in capsys.readouterr().out.split())
    assert printed == ["x_matrix.csv", "x_matrix.json"]


def test_bad_format_exits_2(sample_pdf, capsys):
    with pytest.raises(SystemExit) as e:
        main([str(sample_pdf), "--format", "pdf"])
    assert e.value.code == 2
    assert "formats must be" in capsys.readouterr().err


def test_missing_file_exit_1(tmp_path, capsys):
    assert main([str(tmp_path / "missing.pdf")]) == 1
    assert "not found" in capsys.readouterr().err


def test_scanned_pdf_exit_2(make_pdf, capsys):
    assert main([str(make_pdf([[], []]))]) == 2
    assert "OCR" in capsys.readouterr().err


def test_no_requirements_warns(make_txt, tmp_path, capsys):
    assert main([str(make_txt("Nothing to see here.")), "-o", str(tmp_path)]) == 0
    assert "warning: no shall/must/will/required" in capsys.readouterr().out


def test_summary_uses_singular_for_one(make_txt, tmp_path, capsys):
    text = "SECTION C - PWS\nThe Contractor shall comply with FAR 52.222-41.\n"
    assert main([str(make_txt(text)), "-o", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "(1 page)" in out
    assert "1 requirement: 1 contractor" in out
    assert "1 FAR/DFARS clause referenced" in out


def test_update_sample(answered_xlsx, amendment_pdf, tmp_path, capsys):
    assert main(["update", str(answered_xlsx), str(amendment_pdf), "-o", str(tmp_path), "--format", "xlsx"]) == 0
    out = capsys.readouterr().out
    assert "29 unchanged, 3 changed, 1 new, 1 removed; responses carried to 27 rows" in out
    wb = load_workbook(tmp_path / "sample_solicitation_amd0001_matrix.xlsx")
    assert "Removed" in wb.sheetnames
    ws = wb["Matrix"]
    col = {c.value: i for i, c in enumerate(ws[1])}
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    new = [r for r in rows if r[col["Change"]] == "New"]
    assert len(new) == 1 and "CompTIA A+" in new[0][col["Requirement"]]
    changed = [r for r in rows if r[col["Change"]] == "Changed"]
    due = [r for r in changed if "December 4, 2026" in r[col["Requirement"]]][0]
    assert "November 20, 2026" in due[col["Previous Text"]]
    assert due[col["Compliance"]] == "Comply"
    removed = [" ".join(str(v) for v in r) for r in wb["Removed"].iter_rows(min_row=2, values_only=True)]
    assert any("Resumes are required" in r for r in removed)


def test_update_missing_old(amendment_pdf, tmp_path, capsys):
    assert main(["update", str(tmp_path / "old.xlsx"), str(amendment_pdf), "-o", str(tmp_path)]) == 1
    assert "not found" in capsys.readouterr().err


def test_check_sample(answered_xlsx, capsys):
    assert main(["check", str(answered_xlsx)]) == 3
    out = capsys.readouterr().out
    assert "29 rows to answer (4 Government rows skipped)" in out
    assert "28 answered, 1 blank Compliance" in out
    assert "C.3.2-03" in out and "C.4-02" in out


def test_check_require_extra_column(answered_xlsx, capsys):
    assert main(["check", str(answered_xlsx), "--require", "Compliance,Owner"]) == 3
    assert "blank Owner" in capsys.readouterr().out


def test_check_bad_column_exit_1(answered_xlsx, capsys):
    assert main(["check", str(answered_xlsx), "--require", "Nope"]) == 1
    assert "no column named 'Nope'" in capsys.readouterr().err
