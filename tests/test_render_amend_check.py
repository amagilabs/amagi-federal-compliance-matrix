import csv
import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from amagi_federal_compliance_matrix.amend import carry_forward, normalize
from amagi_federal_compliance_matrix.check import check, report
from amagi_federal_compliance_matrix.detect import Requirement
from amagi_federal_compliance_matrix.render import Matrix, read_matrix_xlsx, to_csv, to_json, to_markdown, write_outputs


def req(id_, text, actor="Contractor", page_end=1):
    return Requirement(
        id=id_, section=id_.split("-")[0], section_title="T", ucf=id_[0], page_start=1, page_end=page_end,
        keyword="shall", kind="Government" if actor == "Government" else "Requirement", actor=actor, text=text,
    )


def matrix(reqs):
    return Matrix(source="x.pdf", pages=2, sha256="0" * 64, requirements=reqs, clauses=[])


def test_markdown_escapes_pipes():
    md = to_markdown(matrix([req("C.1-01", "Use a | b format.")]))
    assert "Use a \\| b format." in md
    assert any(line.startswith("| C.1-01") for line in md.splitlines())


def test_csv_and_json():
    m = matrix([req("C.1-01", "Do it.", page_end=2)])
    rows = list(csv.reader(io.StringIO(to_csv(m).lstrip("\ufeff"))))
    assert rows[0][:3] == ["ID", "Section", "Section Title"]
    assert rows[1][0] == "C.1-01" and "1-2" in rows[1]
    data = json.loads(to_json(m))
    assert data["requirements"][0]["page"] == "1-2"


def test_xlsx_sheets_and_dropdown(tmp_path):
    m = matrix([req("C.1-01", "Do it."), req("M.1-01", "We will rate it.", actor="Government")])
    [path] = write_outputs(m, tmp_path, "t", ["xlsx"])
    assert path.name == "t_matrix.xlsx"
    wb = load_workbook(path)
    assert wb.sheetnames[:2] == ["Matrix", "Clauses"]
    ws = wb["Matrix"]
    header = [c.value for c in ws[1]]
    assert "Compliance" in header and "Response / Notes" in header
    assert ws.freeze_panes == "B2"  # header row and ID column
    assert any("Comply" in (dv.formula1 or "") for dv in ws.data_validations.dataValidation)


def test_read_matrix_requires_columns(tmp_path):
    wb = Workbook()
    wb.active.title = "Matrix"
    wb.active.append(["Foo", "Bar"])
    p = tmp_path / "m.xlsx"
    wb.save(p)
    with pytest.raises(ValueError, match="ID"):
        read_matrix_xlsx(p)


def test_normalize_drops_markers_and_punctuation():
    assert normalize("The Contractor shall: (a) answer CALLS;") == "the contractor shall answer calls"


def _old_workbook(tmp_path, reqs, answers):
    m = matrix(reqs)
    for r in m.requirements:
        r.responses = {"Compliance": answers.get(r.id, "")}
    [p] = write_outputs(m, tmp_path, "old", ["xlsx"])
    wb = load_workbook(p)
    ws = wb["Matrix"]
    col = ws.max_column + 1
    ws.cell(row=1, column=col, value="Writer")  # a column the user added
    ws.cell(row=2, column=col, value="Pat")
    wb.save(p)
    return p


def test_carry_forward(tmp_path):
    old = _old_workbook(tmp_path, [
        req("C.1-01", "The Contractor shall answer calls within 60 seconds."),
        req("C.1-02", "The Contractor shall keep a log of every call."),
        req("C.1-03", "The Contractor shall provide resumes for key staff."),
    ], {"C.1-01": "Comply", "C.1-02": "Partial", "C.1-03": "Comply"})
    new = matrix([
        req("C.1-01", "The Contractor shall answer calls within 45 seconds."),
        req("C.1-02", "The Contractor shall keep a log of every call."),
        req("C.1-03", "The Contractor shall hold an A+ certification."),
    ])
    out = carry_forward(new, old)
    assert [r.change for r in out.requirements] == ["Changed", "Unchanged", "New"]
    first = out.requirements[0]
    assert first.responses["Compliance"] == "Comply"
    assert first.responses["Writer"] == "Pat"
    assert first.previous_text.endswith("60 seconds.")
    assert out.requirements[1].responses["Compliance"] == "Partial"
    assert out.requirements[2].responses["Compliance"] == ""
    assert [r.id for r in out.removed] == ["C.1-03"]
    assert out.extra_columns == ["Writer"]
    assert "Change" in out.columns() and "Previous Text" in out.columns()


def test_carry_forward_threshold(tmp_path):
    old = _old_workbook(tmp_path, [req("C.1-01", "The Contractor shall answer calls within 60 seconds.")],
                        {"C.1-01": "Comply"})
    new = matrix([req("C.1-01", "The Contractor shall answer calls within 45 seconds.")])
    assert carry_forward(new, old, threshold=0.99).requirements[0].change == "New"


def test_check_counts(tmp_path):
    p = _old_workbook(tmp_path, [
        req("C.1-01", "A."), req("C.1-02", "B."), req("C.1-03", "C."), req("C.1-04", "D."),
        req("M.1-01", "E.", actor="Government"),
    ], {"C.1-01": "Comply", "C.1-02": "exception", "C.1-03": "Maybe"})
    res = check(p)
    assert res.considered == 4 and res.skipped_government == 1
    assert [r["ID"] for r in res.blank["Compliance"]] == ["C.1-04"]
    assert [r["ID"] for r in res.exceptions] == ["C.1-02"]
    assert [r["ID"] for r in res.invalid] == ["C.1-03"]
    assert not res.ok
    text = report(res)
    assert "FAIL" in text and "C.1-04" in text
    assert check(p, include_government=True).considered == 5


def test_check_pass_and_missing_column(tmp_path):
    p = _old_workbook(tmp_path, [req("C.1-01", "A.")], {"C.1-01": "N/A"})
    res = check(p)
    assert res.ok and "PASS" in report(res)
    with pytest.raises(ValueError, match="Reviewer"):
        check(p, require=["Reviewer"])
